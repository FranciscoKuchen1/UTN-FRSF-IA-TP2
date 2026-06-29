"""Parse and ingest PDF, DOCX and TXT documents into Supabase pgvector."""

import argparse
import os
import re
from pathlib import Path

from dotenv import load_dotenv
from supabase import Client, create_client

from rag.embeddings import embed_text

load_dotenv(
    dotenv_path=os.path.join(os.path.dirname(__file__), "..", ".env"),
    override=True,
)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}
_supabase: Client | None = None


def _get_supabase() -> Client:
    global _supabase
    if _supabase is None:
        url = os.getenv("SUPABASE_URL", "").strip()
        key = os.getenv("SUPABASE_KEY", "").strip()
        if not url or not key:
            raise RuntimeError("SUPABASE_URL y SUPABASE_KEY deben estar configuradas.")
        _supabase = create_client(url, key)
    return _supabase


def _clean_text(text: str) -> str:
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return "\n".join(line.strip() for line in text.splitlines()).strip()


def parse_pdf(path: str | Path) -> str:
    """Extract selectable text from every page of a PDF."""
    import fitz

    document = fitz.open(path)
    try:
        pages = [_clean_text(page.get_text("text")) for page in document]
    finally:
        document.close()
    return "\n\n".join(page for page in pages if page)


def parse_word(path: str | Path) -> str:
    """Extract non-empty paragraphs and table cells from a DOCX file."""
    from docx import Document

    document = Document(path)
    blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    return "\n\n".join(blocks)


def parse_txt(path: str | Path) -> str:
    """Read plain text, accepting UTF-8 BOM and common Windows encoding."""
    raw = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("El archivo TXT no usa una codificacion compatible (UTF-8 o CP1252).")


def parse_file(path: str | Path) -> str:
    """Validate a file and dispatch it to the matching parser."""
    file_path = Path(path)
    extension = file_path.suffix.lower()
    parsers = {".pdf": parse_pdf, ".docx": parse_word, ".txt": parse_txt}
    if extension not in parsers:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(
            f"Extension '{extension or '<sin extension>'}' no soportada. "
            f"Formatos validos: {supported}"
        )
    if not file_path.is_file():
        raise FileNotFoundError(f"No existe el archivo: {file_path}")

    text = _clean_text(parsers[extension](file_path))
    if not text:
        hint = " El PDF puede ser una imagen escaneada y requerir OCR." if extension == ".pdf" else ""
        raise ValueError(f"No se pudo extraer texto del archivo.{hint}")
    return text


def chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> list[str]:
    """Split text into word-based chunks with a validated overlap."""
    if chunk_size <= 0:
        raise ValueError("chunk_size debe ser mayor que cero.")
    if overlap < 0 or overlap >= chunk_size:
        raise ValueError("overlap debe ser mayor o igual a cero y menor que chunk_size.")

    words = text.split()
    step = chunk_size - overlap
    return [
        " ".join(words[index:index + chunk_size])
        for index in range(0, len(words), step)
        if words[index:index + chunk_size]
    ]


def ingest_document(file_path: str | Path, source_name: str) -> dict:
    """Parse, chunk, embed and persist one supported document."""
    print(f"[INGEST] Parseando '{file_path}'...")
    text = parse_file(file_path)
    chunks = chunk_text(text)
    if not chunks:
        raise ValueError("El documento no genero fragmentos para indexar.")

    print(f"[INGEST] {len(chunks)} chunks generados de '{source_name}'")
    client = _get_supabase()
    for index, chunk in enumerate(chunks):
        embedding = embed_text(chunk)
        if not embedding:
            raise RuntimeError(f"No se genero el embedding del fragmento {index}.")
        client.table("documentos").insert(
            {
                "content": chunk,
                "source": source_name,
                "chunk_index": index,
                "embedding": embedding,
            }
        ).execute()
        print(f"  [{index + 1}/{len(chunks)}] inserted")

    return {
        "source": source_name,
        "format": Path(file_path).suffix.lower(),
        "chunks": len(chunks),
        "characters": len(text),
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Carga un documento PDF, DOCX o TXT en el indice RAG."
    )
    parser.add_argument("--file", required=True, help="Ruta del archivo")
    parser.add_argument("--source", default=None, help="Nombre de fuente")
    args = parser.parse_args()

    summary = ingest_document(
        args.file,
        args.source or os.path.basename(args.file),
    )
    print(f"[INGEST] Completado: {summary}")
